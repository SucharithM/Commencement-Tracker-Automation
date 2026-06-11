from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models.student_details import StudentRecord


HEADERS = [
    "Source Location",
    "Card ID",
    "First Name",
    "Last Name",
    "College",
    "Session",
    "Status",
    "Notes",
]


def export_records_to_excel(records: list[StudentRecord], file_path: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Verified QR Data"

    worksheet.append(HEADERS)
    _style_header(worksheet)

    for record in records:
        worksheet.append(
            [
                record.source_location,
                record.card_id,
                record.first_name,
                record.last_name,
                record.college,
                record.session,
                record.status.value,
                record.error_message,
            ]
        )

    _autosize_columns(worksheet)
    workbook.save(file_path)


def _style_header(worksheet) -> None:
    fill = PatternFill(fill_type="solid", fgColor="1257D1")
    font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def _autosize_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        worksheet.column_dimensions[column_letter].width = min(max_length + 3, 48)
